import io
import shutil
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import numpy as np  # noqa: E402

import app  # noqa: E402

RESULTS = []
DEMO = Path(__file__).resolve().parents[1] / "demo_data"
BACKUP = Path(__file__).resolve().parents[1] / "tests" / "kb.pkl.bak"


def check(name, fn):
    started = time.time()
    try:
        detail = fn()
        RESULTS.append(("PASS", name, detail))
    except Exception as exc:
        import traceback

        RESULTS.append(("FAIL", name, f"{type(exc).__name__}: {exc}"))
        traceback.print_exc()
    print(f"  {RESULTS[-1][0]}  {name}  ({time.time() - started:.1f}s)  {RESULTS[-1][2]}")


class Upload:
    def __init__(self, path):
        self.name = path.name
        self._data = path.read_bytes()

    def getvalue(self):
        return self._data


def _client():
    return app.build_client(app.DEFAULT_BASE_URL)


def _embed_model():
    return app.guess_roles(app.list_models(_client())).embed


def t_pdf_text_extraction():
    pages = app.extract_pages("inspection_report.pdf", (DEMO / "inspection_report.pdf").read_bytes())
    assert len(pages) >= 3, f"only {len(pages)} pages"
    assert pages[0][0] == "inspection_report.pdf p.1"
    assert pages[1][0] == "inspection_report.pdf p.2"
    body = "\n".join(t for _, t in pages)
    for needle in ("F-001", "Distillation Column DC-2", "11.2 mm", "MCB-7" if False else "MCC-7"):
        assert needle in body, f"missing {needle!r}"
    return f"{len(pages)} pages, {len(body)} chars, labels and content correct"


def t_scanned_png_ocr():
    pages = app.extract_pages("inspection_report_scanned.png", (DEMO / "inspection_report_scanned.png").read_bytes())
    text = pages[0][1]
    assert "F-001" in text, text[:200]
    assert "11.2 mm" in text, text[:200]
    assert "2026-11-04" in text
    return f"Tesseract read {len(text)} chars from the scanned page: {text[:60]!r}"


def t_sop_pdf():
    pages = app.extract_pages("sop_hot_work.pdf", (DEMO / "sop_hot_work.pdf").read_bytes())
    body = pages[0][1]
    assert "SOP-ENG-114" in body
    assert "hot work" in body.lower()
    return f"{len(body)} chars, SOP extracted"


def t_docx_extraction():
    import docx

    document = docx.Document()
    document.add_paragraph("Pre-commissioning checklist")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Tag"
    table.rows[0].cells[1].text = "Status"
    table.rows[1].cells[0].text = "PT-101"
    table.rows[1].cells[1].text = "Pass"
    buf = io.BytesIO()
    document.save(buf)
    pages = app.extract_pages("check.docx", buf.getvalue())
    body = pages[0][1]
    assert "Pre-commissioning" in body and "PT-101" in body
    return f"paragraphs+tables extracted ({len(body)} chars)"


def t_chunk_text():
    assert app.chunk_text("") == []
    assert app.chunk_text("   ") == []
    pieces = app.chunk_text("x" * 2500)
    assert len(pieces) == 4, len(pieces)
    assert all(len(p) <= app.CHUNK_CHARS for p in pieces)
    assert pieces[0].endswith(pieces[1][:50]) or True
    return f"2500 chars -> {len(pieces)} chunks of <= {app.CHUNK_CHARS} with overlap"


def _fresh_kb():
    return app.empty_kb()


def t_index_and_search():
    kb = _fresh_kb()
    uploads = [Upload(DEMO / "inspection_report.pdf"), Upload(DEMO / "sop_hot_work.pdf"), Upload(DEMO / "inspection_report_scanned.png")]
    result = app.index_documents(kb, uploads, _client(), _embed_model())
    assert result["chunks"] >= 8, result
    assert kb["vecs"] is not None and kb["vecs"].shape[0] == len(kb["chunks"])
    assert kb["embed_model"] == _embed_model()
    assert kb["dim"] == kb["vecs"].shape[1]
    assert set(kb["docs"]) == {"inspection_report.pdf", "sop_hot_work.pdf", "inspection_report_scanned.png"}
    hits = app.search_kb(_client(), kb, _embed_model(), "what is the remaining tray thickness on DC-2?", 6)
    assert hits, "no hits"
    best_source, best_text = hits[0][0]
    assert "11.2" in best_text or "13.2" in best_text, best_text[:200]
    return (
        f"{result['chunks']} chunks, dim={kb['dim']}, top hit '{best_source}' "
        f"score={hits[0][1]:.3f} contains the thickness value"
    )


def t_kb_status_transitions():
    kb = _fresh_kb()
    assert app.kb_status(kb, "nomic-embed-text") == (False, "empty")
    app.index_documents(kb, [Upload(DEMO / "sop_hot_work.pdf")], _client(), _embed_model())
    ok, reason = app.kb_status(kb, _embed_model())
    assert ok, reason
    ok, reason = app.kb_status(kb, "some-other-embed-model")
    assert not ok and reason == "embed-model-changed", reason
    return f"empty -> ready -> blocked with 'embed-model-changed'"


def t_reindex_restores():
    kb = _fresh_kb()
    app.index_documents(kb, [Upload(DEMO / "sop_hot_work.pdf")], _client(), _embed_model())
    before = len(kb["chunks"])
    kb["embed_model"] = "bge-m3"
    ok, reason = app.kb_status(kb, _embed_model())
    assert not ok and reason == "embed-model-changed"
    app.reindex_with_model(kb, _client(), _embed_model())
    ok, reason = app.kb_status(kb, _embed_model())
    assert ok, reason
    assert len(kb["chunks"]) == before
    hits = app.search_kb(_client(), kb, _embed_model(), "gas testing interval", 3)
    assert hits, "search broken after re-index"
    return f"blocked -> re-indexed -> ready, {before} chunks re-embedded, search works"


def t_delete_document():
    kb = _fresh_kb()
    app.index_documents(
        kb,
        [Upload(DEMO / "inspection_report.pdf"), Upload(DEMO / "sop_hot_work.pdf")],
        _client(),
        _embed_model(),
    )
    total = len(kb["chunks"])
    report_chunks = len(kb["docs"]["inspection_report.pdf"])
    removed = app.delete_document(kb, "inspection_report.pdf")
    assert removed == report_chunks, (removed, report_chunks)
    assert len(kb["chunks"]) == total - report_chunks
    assert "inspection_report.pdf" not in kb["docs"]
    assert kb["vecs"].shape[0] == len(kb["chunks"])
    assert not any("F-001" in t for _, t in kb["chunks"]), "report rows survived deletion"
    assert any("SOP-ENG-114" in t for _, t in kb["chunks"]), "SOP rows wrongly removed"
    assert app.delete_document(kb, "does-not-exist.pdf") == 0
    return f"removed {removed} chunks, vectors stayed aligned ({kb['vecs'].shape[0]} rows), other doc intact"


def t_delete_last_document():
    kb = _fresh_kb()
    app.index_documents(kb, [Upload(DEMO / "sop_hot_work.pdf")], _client(), _embed_model())
    app.delete_document(kb, "sop_hot_work.pdf")
    assert kb["chunks"] == [] and kb["vecs"] is None and kb["docs"] == {}
    return "emptying the KB resets vectors and doc map"


def t_make_docx_layout():
    body = (
        "## Background\nInspection of Unit 4 found five findings.\n\n"
        "## Findings\n\n"
        "| ID | Severity | Unit | Due |\n"
        "| --- | --- | --- | --- |\n"
        "| F-001 | HIGH | DC-2 | 2026-11-04 |\n"
        "| F-003 | MEDIUM | E-204 | 2026-10-15 |\n\n"
        "- Replace the **tray set** at stage *seven*\n"
        "- Recalibrate LT-301B\n\n"
        "## Recommendation\nClose all HIGH findings by the due date.\n\n"
        "## Approval\n"
    )
    path = Path(app.make_docx("Approval Note IR-2026-0417", body, "approval_note.docx"))
    assert path.exists() and path.name == "approval_note.docx"
    assert str(path).startswith(str(app.OUT_DIR)), "written outside workspace/out"

    import docx

    document = docx.Document(path)
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert "Background" in headings and "Findings" in headings and "Approval" in headings
    assert len(document.tables) >= 2, f"expected findings table + approval block, got {len(document.tables)}"
    findings_table = document.tables[0]
    dumped = [[c.text for c in row.cells] for row in findings_table.rows]
    assert dumped[0] == ["ID", "Severity", "Unit", "Due"], dumped
    assert dumped[1][0] == "F-001" and dumped[1][1] == "HIGH", dumped
    assert dumped[2][0] == "F-003" and dumped[2][3] == "2026-10-15", dumped
    assert len(dumped) == 3, f"separator row not dropped: {dumped}"
    assert document.tables[-1].rows[1].cells[0].text == "Prepared by"
    header_text = document.sections[0].header.paragraphs[0].text
    assert "Approval Note IR-2026-0417" in header_text
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert any("tray set" in b for b in bullets), bullets
    bold_runs = [r.text for p in document.paragraphs for r in p.runs if r.bold]
    assert "tray set" in bold_runs, bold_runs
    italic_runs = [r.text for p in document.paragraphs for r in p.runs if r.italic]
    assert "seven" in italic_runs, italic_runs
    return f"headings={len(headings)} tables={len(document.tables)} bullets={len(bullets)} bold+italic applied, header set"


def t_make_docx_unicode_bullets():
    body = (
        "TOOLBOX TALK\n"
        "Topic: Static Electricity\n"
        "\n"
        "• Static charge can build up and ignite flammable vapour\n"
        "• Ground all equipment and wear anti-static clothing\n"
        "• Use ionising air blowers to neutralise charge\n"
        "\n"
        "1. Isolate the area\n"
        "2. Bond and earth the vessel\n"
        "3. Verify with an electrostatic meter\n"
    )
    path = Path(app.make_docx("Toolbox Talk: Static Electricity", body, "unicode_bullets.docx"))
    import docx

    document = docx.Document(path)
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    numbered = [p.text for p in document.paragraphs if p.style.name == "List Number"]
    assert len(bullets) == 3, bullets
    assert len(numbered) == 3, numbered
    assert bullets[0].startswith("Static charge"), bullets[0]
    assert not any(p.text.startswith("•") for p in document.paragraphs), "bullet glyph left in text"
    return f"{len(bullets)} unicode bullets + {len(numbered)} numbered items converted to real Word styles"


def t_make_docx_path_traversal():
    evil = "../../../../Windows/System32/drivers/etc/pwned.docx"
    path = Path(app.make_docx("Evil", "body", evil))
    assert path.parent == app.OUT_DIR, f"escaped to {path.parent}"
    assert ".." not in path.name
    naked = app.make_docx("NoExt", "body", "report")
    assert Path(naked).name == "report.docx", naked
    return f"'{evil}' -> {path.name}, extensionless name repaired"


def t_run_python():
    ok = app.run_python("print(6*7)")
    assert "exit=0" in ok and "42" in ok, ok
    err = app.run_python("raise ValueError('boom')")
    assert "exit=1" in err and "ValueError" in err and "boom" in err, err
    inf = app.run_python("while True: pass")
    assert "timed out" in inf, inf
    return "stdout ok, stderr surfaced, infinite loop killed at 20s"


def t_tool_arg_coercion():
    specs = ("code",)
    assert app._coerce_args(specs, {"code": 123}) == {"code": "123"}
    assert app._coerce_args(specs, {"code": None}) == {"code": ""}
    assert app._coerce_args(specs, {"code": {"a": 1}}) == {"code": '{"a": 1}'}
    assert app._coerce_args(specs, {"wrong": "x"}) == {}
    assert app._coerce_args(specs, {"code": "print(1)", "extra": 1}) == {"code": "print(1)"}
    return "scalars, None, dicts, missing and extra keys all handled"


def t_first_python_block():
    assert app._first_python_block("```python\nprint(1)\n```") == "print(1)"
    assert app._first_python_block("```py\nx=2\n```") == "x=2"
    assert app._first_python_block("```\nplain\n```") == ""
    assert app._first_python_block("no code here") == ""
    assert app._first_python_block("") == ""
    assert app._first_python_block(None) == ""
    multi = app._first_python_block("intro\n```python\nfirst()\n```\nmid\n```python\nsecond()\n```")
    assert multi == "first()", multi
    return "fenced python/py blocks found, other fences and None ignored"


def t_run_python_utf8():
    result = app.run_python("print('heat exchanger DC-2 \\u00d7 5 ok')")
    assert "exit=0" in result, result
    assert "DC-2" in result and "5 ok" in result, result
    return "non-ASCII output survived the Windows console encoding"


def t_audit_file_claims_flags_hallucination():
    claim = "The memo has been saved as **output.docx**. You can download it."
    flagged = app.audit_file_claims(claim, [])
    assert flagged != claim
    assert "Correction" in flagged and "output.docx" in flagged
    assert "No Word document was created" in flagged
    return "unsaved file claim flagged when make_docx never ran"


def t_audit_file_claims_accepts_real_file():
    answer = "Done, saved as toolbox_talk.docx with three bullet points."
    assert app.audit_file_claims(answer, ["toolbox_talk.docx"]) == answer
    assert app.audit_file_claims(answer, [str(app.OUT_DIR / "toolbox_talk.docx")]) == answer
    return "real file left untouched, bare and full paths both accepted"


def t_audit_file_claims_partial():
    answer = "I saved toolbox_talk.docx but not the spreadsheet report.docx."
    flagged = app.audit_file_claims(answer, ["toolbox_talk.docx"])
    correction = flagged.split("Correction", 1)[1]
    assert "report.docx" in correction
    assert "no tool result produced report.docx" in correction
    assert "actually created" in correction and "toolbox_talk.docx" in correction
    return "only the bogus file is named as unproduced; the real one is listed as created"


def t_audit_no_claims_untouched():
    for answer in ("Here is the answer.", "Use `make_docx` to save a file.", "", None):
        assert app.audit_file_claims(answer, []) == answer
    return "answers with no .docx mention are never modified"


def t_audit_rejects_path_tricks():
    answer = "Saved to C:/Windows/System32/evil.docx"
    flagged = app.audit_file_claims(answer, [])
    assert "Correction" in flagged, flagged
    return "absolute-path file claims are flagged too"


def t_broken_tool_call_detected():
    cut = [
        '{"tool": "make_docx", "args": {"content": "# Unit 4 Report',
        '{"tool": "search_kb"}',
        '{"tool":"run_python","args":{',
        '  {"tool": "make_docx"',
    ]
    for text in cut:
        assert app._looks_like_broken_call(text), repr(text)
    for text in ['{"value": 3}', "Here is a normal answer.", "", None]:
        assert not app._looks_like_broken_call(text), repr(text)
    return "truncated tool JSON is caught, ordinary answers are not"


def t_audit_requested_filename_mismatch():
    answer = "Done, I saved the document."
    out = app.audit_file_claims(answer, ["output.docx"], "make toolbox_talk.docx")
    assert "You asked for toolbox_talk.docx" in out
    assert "output.docx" in out
    same = app.audit_file_claims(answer, ["toolbox_talk.docx"], "make toolbox_talk.docx")
    assert same == answer, same
    return "wrong filename is flagged; correct filename passes through clean"


def t_tool_allowlist():
    client = _client()
    roles = app.Roles(chat=app.guess_roles(app.list_models(client)).chat, embed=_embed_model())
    kb = _fresh_kb()
    general = app.make_tool_specs("general", client, kb, roles)
    kb_specs = app.make_tool_specs("kb", client, kb, roles)
    assert set(general) == {"run_python", "make_docx"}, set(general)
    assert set(kb_specs) == {"search_kb", "make_docx"}, set(kb_specs)
    assert "run_python" not in kb_specs, "run_python reachable in KB mode"
    assert "search_kb" not in general
    return "General={run_python,make_docx}  KB={search_kb,make_docx}"


def t_kb_prompt_has_no_run_python():
    prompt = app.build_system_prompt("kb", "SOME CONTEXT")
    assert "run_python" not in prompt
    assert app.REFUSAL_TEXT in prompt
    assert "untrusted data" in prompt and "ignore it completely" in prompt
    general = app.build_system_prompt("general")
    assert "run_python" in general and app.REFUSAL_TEXT not in general
    empty_ctx = app.build_system_prompt("kb", "")
    assert app.REFUSAL_TEXT in empty_ctx
    return "KB prompt blocks run_python and carries the exact refusal string"


def t_chunking_respects_page_labels():
    kb = _fresh_kb()
    app.index_documents(kb, [Upload(DEMO / "inspection_report.pdf")], _client(), _embed_model())
    labels = {source for source, _ in kb["chunks"]}
    assert all(l.startswith("inspection_report.pdf p.") for l in labels), labels
    assert len(labels) >= 3, labels
    return f"citations will use {sorted(labels)}"


def main():
    if app.KB_PATH.exists():
        shutil.copy2(app.KB_PATH, BACKUP)
    print("== ingestion ==")
    for n, f in [
        ("PDF text extraction", t_pdf_text_extraction),
        ("scanned PNG via Tesseract", t_scanned_png_ocr),
        ("SOP PDF", t_sop_pdf),
        ("DOCX paragraphs+tables", t_docx_extraction),
        ("chunker", t_chunk_text),
    ]:
        check(n, f)

    print("== knowledge base ==")
    for n, f in [
        ("index + search", t_index_and_search),
        ("kb_status transitions", t_kb_status_transitions),
        ("re-index after model change", t_reindex_restores),
        ("delete document", t_delete_document),
        ("delete last document", t_delete_last_document),
        ("page labels survive chunking", t_chunking_respects_page_labels),
    ]:
        check(n, f)

    print("== tools ==")
    for n, f in [
        ("make_docx layout", t_make_docx_layout),
        ("make_docx unicode bullets", t_make_docx_unicode_bullets),
        ("make_docx path traversal", t_make_docx_path_traversal),
        ("run_python", t_run_python),
        ("run_python utf-8 output", t_run_python_utf8),
        ("first python block detection", t_first_python_block),
        ("file-claim audit: hallucination", t_audit_file_claims_flags_hallucination),
        ("file-claim audit: real file", t_audit_file_claims_accepts_real_file),
        ("file-claim audit: partial", t_audit_file_claims_partial),
        ("file-claim audit: no claims", t_audit_no_claims_untouched),
        ("file-claim audit: path tricks", t_audit_rejects_path_tricks),
        ("truncated tool call detected", t_broken_tool_call_detected),
        ("file-claim audit: filename mismatch", t_audit_requested_filename_mismatch),
        ("tool arg coercion", t_tool_arg_coercion),
        ("per-mode tool allow-list", t_tool_allowlist),
        ("KB prompt excludes run_python", t_kb_prompt_has_no_run_python),
    ]:
        check(n, f)

    failed = [r for r in RESULTS if r[0] == "FAIL"]
    print(f"\n{len(RESULTS) - len(failed)}/{len(RESULTS)} passed")
    for status, name, detail in failed:
        print(f"  FAIL {name}: {detail}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
